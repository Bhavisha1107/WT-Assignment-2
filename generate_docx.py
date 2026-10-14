import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_element(name):
    return OxmlElement(name)

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_code_block(doc, code_text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    # Border
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
            <w:left w:val="single" w:sz="18" w:space="0" w:color="3B82F6"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    
    run = p.add_run(code_text.strip())
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(30, 41, 59)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def read_file(path):
    with open(path, 'r', encoding='utf-8') as f:
        return f.read()

def main():
    base_dir = r"d:\Subject\WT\Assignment-2"
    doc = Document()
    
    # Page setup - 0.7 inch margins
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(0.7)
        s.bottom_margin = Inches(0.7)
        s.left_margin = Inches(0.7)
        s.right_margin = Inches(0.7)
    
    # Header Banner Table
    banner_table = doc.add_table(rows=1, cols=1)
    banner_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    bcell = banner_table.cell(0, 0)
    bcell.width = Inches(6.8)
    set_cell_background(bcell, "1E3A8A")
    set_cell_margins(bcell, top=200, bottom=200, left=250, right=250)
    
    bp = bcell.paragraphs[0]
    bp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = bp.add_run("ASSIGNMENT – 2\n")
    r1.font.name = "Arial"
    r1.font.size = Pt(20)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(255, 255, 255)
    
    r2 = bp.add_run("U21CS501 – WEB TECHNOLOGY\n")
    r2.font.name = "Arial"
    r2.font.size = Pt(14)
    r2.font.bold = True
    r2.font.color.rgb = RGBColor(191, 219, 254)
    
    r3 = bp.add_run("Department of Computer Science and Engineering")
    r3.font.name = "Arial"
    r3.font.size = Pt(11)
    r3.font.color.rgb = RGBColor(224, 231, 255)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    
    # Student Info & GitHub Box
    meta_table = doc.add_table(rows=2, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r in meta_table.rows:
        for c in r.cells:
            c.width = Inches(3.4)
            set_cell_background(c, "F1F5F9")
            set_cell_margins(c, top=80, bottom=80, left=150, right=150)
            
    p00 = meta_table.cell(0, 0).paragraphs[0]
    p00.add_run("Student Name: ").bold = True
    p00.add_run("Bhavisha S")
    
    p01 = meta_table.cell(0, 1).paragraphs[0]
    p01.add_run("Roll Number: ").bold = True
    p01.add_run("24CS031")
    
    p10 = meta_table.cell(1, 0).paragraphs[0]
    p10.add_run("Academic Course: ").bold = True
    p10.add_run("Web Technology (U21CS501)")
    
    p11 = meta_table.cell(1, 1).paragraphs[0]
    p11.add_run("GitHub Link: ").bold = True
    gh_run = p11.add_run("https://github.com/Bhavisha1107/WT-Assignment-2")
    gh_run.font.color.rgb = RGBColor(37, 99, 235)
    gh_run.font.underline = True
    
    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    
    # -------------------------------------------------------------
    # QUESTION 1
    # -------------------------------------------------------------
    h1 = doc.add_heading("Question 1: Student Registration Form with Validation", level=1)
    h1.runs[0].font.color.rgb = RGBColor(30, 58, 138)
    
    p_sc1 = doc.add_paragraph()
    r = p_sc1.add_run("Scenario: ")
    r.bold = True
    p_sc1.add_run("The student admission portal requires a well-designed registration form with proper validation and styling.")
    
    p_tk1 = doc.add_paragraph()
    r = p_tk1.add_run("Tasks:\n")
    r.bold = True
    p_tk1.add_run("i. Create HTML forms for student registration (name, roll number, email, course, file upload).\n")
    p_tk1.add_run("ii. Apply CSS styles for layout, colors, and responsive design (e.g., styled input fields, buttons, error messages).\n")
    p_tk1.add_run("iii. Implement client-side validation (JavaScript) for required fields, email format, and password strength.")
    
    # Codes
    doc.add_heading("1. Index.html", level=2)
    q1_html = read_file(os.path.join(base_dir, "Question-1_Student_Registration", "index.html"))
    add_code_block(doc, q1_html)
    
    doc.add_heading("2. Style.css", level=2)
    q1_css = read_file(os.path.join(base_dir, "Question-1_Student_Registration", "style.css"))
    add_code_block(doc, q1_css)
    
    doc.add_heading("3. Validation.js", level=2)
    q1_js = read_file(os.path.join(base_dir, "Question-1_Student_Registration", "validation.js"))
    add_code_block(doc, q1_js)
    
    doc.add_heading("Output - Question 1:", level=2)
    doc.add_paragraph("The output showcases the responsive registration card with real-time field error indicators, password requirements, and successful validation state:")
    
    q1_empty_img = os.path.join(base_dir, "output_screenshots", "q1_empty.png")
    q1_filled_img = os.path.join(base_dir, "output_screenshots", "q1_filled.png")
    
    if os.path.exists(q1_empty_img) and os.path.exists(q1_filled_img):
        img_table = doc.add_table(rows=1, cols=2)
        img_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        c1, c2 = img_table.cell(0, 0), img_table.cell(0, 1)
        c1.width = Inches(3.2)
        c2.width = Inches(3.2)
        
        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p1.add_run().add_picture(q1_empty_img, width=Inches(3.0))
        p1_cap = c1.add_paragraph("Figure 1.1: Initial Registration Form")
        p1_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p1_cap.runs[0].font.size = Pt(8.5)
        
        p2 = c2.paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.add_run().add_picture(q1_filled_img, width=Inches(3.0))
        p2_cap = c2.add_paragraph("Figure 1.2: Validated Form with Success State")
        p2_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2_cap.runs[0].font.size = Pt(8.5)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    
    p_c1 = doc.add_paragraph()
    r = p_c1.add_run("Conclusion:\n")
    r.bold = True
    p_c1.add_run("Thus, the Student Registration Form was successfully created using HTML, CSS, and JavaScript. HTML provides structured input fields for student admission details; CSS provides an intuitive responsive layout with soft gradient accents, elevation shadows, and styled focus states; JavaScript enforces thorough client-side validation for mandatory fields, valid email pattern matching, strong password criteria, and document size/type restrictions. This improves data accuracy and provides a seamless registration experience.")
    
    doc.add_page_break()
    
    # -------------------------------------------------------------
    # QUESTION 2
    # -------------------------------------------------------------
    h2 = doc.add_heading("Question 2: College Admin Panel with CRUD & Document Upload", level=1)
    h2.runs[0].font.color.rgb = RGBColor(30, 58, 138)
    
    p_sc2 = doc.add_paragraph()
    r = p_sc2.add_run("Scenario: ")
    r.bold = True
    p_sc2.add_run("The college admin panel must manage student records and their profile documents/images.")
    
    p_tk2 = doc.add_paragraph()
    r = p_tk2.add_run("Tasks:\n")
    r.bold = True
    p_tk2.add_run("1. CRUD Operations (Add, Edit, Delete, View) for student details (name, roll number, course, etc.) using MySQL.\n")
    p_tk2.add_run("2. File Upload Handling for profile images and documents (e.g., ID proof, certificates).\n")
    p_tk2.add_run("3. Store Metadata in MySQL — file path, file type, upload date linked to the student record.\n")
    p_tk2.add_run("4. Display Uploaded Files when viewing student details.\n")
    p_tk2.add_run("5. Validation & Error Handling — check file size, type, and prevent overwriting.")
    
    doc.add_heading("1. Database Configuration (db.php)", level=2)
    add_code_block(doc, read_file(os.path.join(base_dir, "Question-2_PHP_Admin_Panel", "db.php")))
    
    doc.add_heading("2. Student Records Dashboard (index.php)", level=2)
    add_code_block(doc, read_file(os.path.join(base_dir, "Question-2_PHP_Admin_Panel", "index.php")))
    
    doc.add_heading("3. Add Student (add.php)", level=2)
    add_code_block(doc, read_file(os.path.join(base_dir, "Question-2_PHP_Admin_Panel", "add.php")))
    
    doc.add_heading("4. Edit Student (edit.php)", level=2)
    add_code_block(doc, read_file(os.path.join(base_dir, "Question-2_PHP_Admin_Panel", "edit.php")))
    
    doc.add_heading("5. Delete Student (delete.php)", level=2)
    add_code_block(doc, read_file(os.path.join(base_dir, "Question-2_PHP_Admin_Panel", "delete.php")))
    
    doc.add_heading("6. Document Upload (upload.php)", level=2)
    add_code_block(doc, read_file(os.path.join(base_dir, "Question-2_PHP_Admin_Panel", "upload.php")))
    
    doc.add_heading("7. View Student & Files (view.php)", level=2)
    add_code_block(doc, read_file(os.path.join(base_dir, "Question-2_PHP_Admin_Panel", "view.php")))
    
    doc.add_heading("8. Database Schema (student_admin.sql)", level=2)
    add_code_block(doc, read_file(os.path.join(base_dir, "Question-2_PHP_Admin_Panel", "student_admin.sql")))
    
    doc.add_heading("Output - Question 2:", level=2)
    doc.add_paragraph("The redesigned PHP Admin Panel displays student admission records with styled actions, validation warnings, and file preview cards:")
    
    q2_idx_img = os.path.join(base_dir, "output_screenshots", "q2_index.png")
    if os.path.exists(q2_idx_img):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(q2_idx_img, width=Inches(5.8))
        cap = doc.add_paragraph("Figure 2.1: Student Admin Panel Dashboard with Records & Actions")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.runs[0].font.size = Pt(8.5)
        
    q2_add_img = os.path.join(base_dir, "output_screenshots", "q2_add.png")
    q2_upl_img = os.path.join(base_dir, "output_screenshots", "q2_upload.png")
    if os.path.exists(q2_add_img) and os.path.exists(q2_upl_img):
        img_table2 = doc.add_table(rows=1, cols=2)
        img_table2.alignment = WD_TABLE_ALIGNMENT.CENTER
        c1, c2 = img_table2.cell(0, 0), img_table2.cell(0, 1)
        c1.width = Inches(3.2)
        c2.width = Inches(3.2)
        
        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p1.add_run().add_picture(q2_add_img, width=Inches(3.0))
        p1_cap = c1.add_paragraph("Figure 2.2: Add Student Form")
        p1_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p1_cap.runs[0].font.size = Pt(8.5)
        
        p2 = c2.paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.add_run().add_picture(q2_upl_img, width=Inches(3.0))
        p2_cap = c2.add_paragraph("Figure 2.3: Document Upload Handling")
        p2_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2_cap.runs[0].font.size = Pt(8.5)
        
    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    
    p_c2 = doc.add_paragraph()
    r = p_c2.add_run("Conclusion:\n")
    r.bold = True
    p_c2.add_run("Thus, the Student Admin Panel was successfully developed using PHP and MySQL to manage student records through complete CRUD operations. File upload functionality was implemented to store profile images and documents securely. File metadata, including storage path, MIME type, upload timestamp, and original file name, is maintained in MySQL with foreign key cascading. Server-side validation, PDO prepared statements, and XSS sanitization ensure robust protection against injection and unauthorized overwriting.")
    
    doc.add_page_break()
    
    # -------------------------------------------------------------
    # QUESTION 3
    # -------------------------------------------------------------
    h3 = doc.add_heading("Question 3: Student Profile Portal with Java Servlets, JSP & JDBC", level=1)
    h3.runs[0].font.color.rgb = RGBColor(30, 58, 138)
    
    p_sc3 = doc.add_paragraph()
    r = p_sc3.add_run("Scenario: ")
    r.bold = True
    p_sc3.add_run("The front-end portal must allow users to view student profiles along with their uploaded images/documents.")
    
    p_tk3 = doc.add_paragraph()
    r = p_tk3.add_run("Tasks:\n")
    r.bold = True
    p_tk3.add_run("1. Connect to MySQL using JDBC to fetch student details and associated file paths.\n")
    p_tk3.add_run("2. Retrieve and stream images/documents from the folder where they are stored.\n")
    p_tk3.add_run("3. Display student profiles dynamically in an HTML/JSP page with <img> tags pointing to the Servlet.\n")
    p_tk3.add_run("4. Implement looping and conditional logic to list multiple students and their files.\n")
    p_tk3.add_run("5. Ensure session handling so that only logged-in users can access the portal securely.")
    
    doc.add_heading("1. Project Object Model (pom.xml)", level=2)
    add_code_block(doc, read_file(os.path.join(base_dir, "Question-3_Java_Servlet_Portal", "pom.xml")))
    
    doc.add_heading("2. Deployment Descriptor (web.xml)", level=2)
    add_code_block(doc, read_file(os.path.join(base_dir, "Question-3_Java_Servlet_Portal", "src", "main", "webapp", "WEB-INF", "web.xml")))
    
    doc.add_heading("3. Database Connection Utility (DBConnection.java)", level=2)
    add_code_block(doc, read_file(os.path.join(base_dir, "Question-3_Java_Servlet_Portal", "src", "main", "java", "com", "college", "DBConnection.java")))
    
    doc.add_heading("4. Student File Entity (StudentFile.java)", level=2)
    add_code_block(doc, read_file(os.path.join(base_dir, "Question-3_Java_Servlet_Portal", "src", "main", "java", "com", "college", "StudentFile.java")))
    
    doc.add_heading("5. Student Entity Model (Student.java)", level=2)
    add_code_block(doc, read_file(os.path.join(base_dir, "Question-3_Java_Servlet_Portal", "src", "main", "java", "com", "college", "Student.java")))
    
    doc.add_heading("6. Authentication Controller (LoginServlet.java)", level=2)
    add_code_block(doc, read_file(os.path.join(base_dir, "Question-3_Java_Servlet_Portal", "src", "main", "java", "com", "college", "LoginServlet.java")))
    
    doc.add_heading("7. Session Invalidation Controller (LogoutServlet.java)", level=2)
    add_code_block(doc, read_file(os.path.join(base_dir, "Question-3_Java_Servlet_Portal", "src", "main", "java", "com", "college", "LogoutServlet.java")))
    
    doc.add_heading("8. Student Profiles Controller (StudentServlet.java)", level=2)
    add_code_block(doc, read_file(os.path.join(base_dir, "Question-3_Java_Servlet_Portal", "src", "main", "java", "com", "college", "StudentServlet.java")))
    
    doc.add_heading("9. Secure Document Streaming (FileServlet.java)", level=2)
    add_code_block(doc, read_file(os.path.join(base_dir, "Question-3_Java_Servlet_Portal", "src", "main", "java", "com", "college", "FileServlet.java")))
    
    doc.add_heading("10. Authentication View (login.jsp)", level=2)
    add_code_block(doc, read_file(os.path.join(base_dir, "Question-3_Java_Servlet_Portal", "src", "main", "webapp", "login.jsp")))
    
    doc.add_heading("11. Dynamic Profiles View (students.jsp)", level=2)
    add_code_block(doc, read_file(os.path.join(base_dir, "Question-3_Java_Servlet_Portal", "src", "main", "webapp", "students.jsp")))
    
    doc.add_heading("12. MySQL Database Schema & Mock Data (database.sql)", level=2)
    add_code_block(doc, read_file(os.path.join(base_dir, "Question-3_Java_Servlet_Portal", "database.sql")))
    
    doc.add_heading("Output - Question 3:", level=2)
    doc.add_paragraph("The Student Profile Portal outputs include a secured login interface, a responsive multi-card student directory, and a dedicated document viewer:")
    
    q3_log_img = os.path.join(base_dir, "output_screenshots", "q3_login.png")
    if os.path.exists(q3_log_img):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(q3_log_img, width=Inches(5.5))
        cap = doc.add_paragraph("Figure 3.1: Student Profile Portal Login Screen")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.runs[0].font.size = Pt(8.5)
        
    q3_stu_img = os.path.join(base_dir, "output_screenshots", "q3_students.png")
    if os.path.exists(q3_stu_img):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(q3_stu_img, width=Inches(5.8))
        cap = doc.add_paragraph("Figure 3.2: Dynamic Student Directory with Profile Cards & Search")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.runs[0].font.size = Pt(8.5)
        
    q3_doc_img = os.path.join(base_dir, "output_screenshots", "q3_documents.png")
    if os.path.exists(q3_doc_img):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(q3_doc_img, width=Inches(5.8))
        cap = doc.add_paragraph("Figure 3.3: Student Document Repository View")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.runs[0].font.size = Pt(8.5)
        
    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    
    p_c3 = doc.add_paragraph()
    r = p_c3.add_run("Conclusion:\n")
    r.bold = True
    p_c3.add_run("The Student Profile Portal was developed using Java Servlets (Jakarta EE 6.0), JSP, JSTL, JDBC, and MySQL. It displays student details and provides secure, streamed access to images and PDF documents directly through FileServlet, preventing direct server directory exposure. Session management safeguards all portal endpoints against unauthorized browsing. The architecture demonstrates MVC separation of concerns, secure file streaming, and responsive presentation.")
    
    # Save the document
    output_filename = os.path.join(base_dir, "Assignment-2_Web_Technology_Bhavisha_S_24CS031.docx")
    doc.save(output_filename)
    print(f"Successfully generated: {output_filename}")

if __name__ == "__main__":
    main()
